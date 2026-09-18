from dotenv import load_dotenv
from pathlib import Path
import os

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

from fastapi import FastAPI, APIRouter, HTTPException, Depends, Request, Response, UploadFile, File, Form, Query, Header
from fastapi.responses import StreamingResponse
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import logging
from pydantic import BaseModel, Field, EmailStr
from typing import List, Optional, Dict, Any
import uuid
import io
import re
from datetime import datetime, timezone, timedelta
import bcrypt
import jwt
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
EMERGENT_KEY = os.environ.get("EMERGENT_LLM_KEY")
APP_NAME = "siscoring-kalteng"
storage_key = None

def init_storage(force: bool = False):
    global storage_key
    if storage_key and not force:
        return storage_key
    resp = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": EMERGENT_KEY}, timeout=30)
    resp.raise_for_status()
    storage_key = resp.json()["storage_key"]
    return storage_key

def put_object(path: str, data: bytes, content_type: str) -> dict:
    key = init_storage()
    resp = requests.put(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key, "Content-Type": content_type}, data=data, timeout=120)
    if resp.status_code == 404:
        key = init_storage(force=True)
        resp = requests.put(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key, "Content-Type": content_type}, data=data, timeout=120)
    resp.raise_for_status()
    return resp.json()

def get_object(path: str):
    key = init_storage()
    resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    if resp.status_code == 404:
        key = init_storage(force=True)
        resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    resp.raise_for_status()
    return resp.content, resp.headers.get("Content-Type", "application/octet-stream")

# ---------------- Auth ----------------
JWT_ALGORITHM = "HS256"
def get_jwt_secret(): return os.environ["JWT_SECRET"]
def hash_password(p): return bcrypt.hashpw(p.encode(), bcrypt.gensalt()).decode()
def verify_password(p, h):
    try: return bcrypt.checkpw(p.encode(), h.encode())
    except Exception: return False
def create_access_token(uid):
    return jwt.encode({"sub": uid, "exp": datetime.now(timezone.utc) + timedelta(days=7), "type": "access"}, get_jwt_secret(), algorithm=JWT_ALGORITHM)
def set_auth_cookie(response, token):
    response.set_cookie("access_token", token, httponly=True, secure=True, samesite="none", max_age=604800, path="/")

async def get_current_user(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        ah = request.headers.get("Authorization", "")
        if ah.startswith("Bearer "): token = ah[7:]
    if not token: raise HTTPException(status_code=401, detail="Tidak terautentikasi")
    try:
        payload = jwt.decode(token, get_jwt_secret(), algorithms=[JWT_ALGORITHM])
        user = await db.users.find_one({"id": payload["sub"]}, {"_id": 0})
        if not user or not user.get("active", True):
            raise HTTPException(status_code=401, detail="Pengguna tidak ditemukan / nonaktif")
        user.pop("password_hash", None)
        return user
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Sesi berakhir")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")

def require_roles(*roles):
    async def checker(user: dict = Depends(get_current_user)):
        if user["role"] not in roles:
            raise HTTPException(status_code=403, detail="Akses ditolak untuk peran ini")
        return user
    return checker

# ---------------- Reference data ----------------
AREAS = [
    {"name": "Provinsi Kalimantan Tengah", "level": "provinsi"},
    {"name": "Kabupaten Barito Selatan", "level": "kabupaten"},
    {"name": "Kabupaten Barito Timur", "level": "kabupaten"},
    {"name": "Kabupaten Barito Utara", "level": "kabupaten"},
    {"name": "Kabupaten Gunung Mas", "level": "kabupaten"},
    {"name": "Kabupaten Kapuas", "level": "kabupaten"},
    {"name": "Kabupaten Katingan", "level": "kabupaten"},
    {"name": "Kabupaten Kotawaringin Barat", "level": "kabupaten"},
    {"name": "Kabupaten Kotawaringin Timur", "level": "kabupaten"},
    {"name": "Kabupaten Lamandau", "level": "kabupaten"},
    {"name": "Kabupaten Murung Raya", "level": "kabupaten"},
    {"name": "Kabupaten Pulang Pisau", "level": "kabupaten"},
    {"name": "Kabupaten Seruyan", "level": "kabupaten"},
    {"name": "Kabupaten Sukamara", "level": "kabupaten"},
    {"name": "Kota Palangka Raya", "level": "kabupaten"},
]
AREA_LEVEL = {a["name"]: a["level"] for a in AREAS}

URUSAN_STATIC = [
    "Bidang Pendidikan", "Bidang Kesehatan", "Bidang Pekerjaan Umum dan Penataan Ruang",
    "Bidang Perumahan dan Kawasan Permukiman",
    "Bidang Ketentraman dan Ketertiban Umum serta Perlindungan Masyarakat",
    "Bidang Sosial", "Bidang Tenaga Kerja", "Bidang Pemberdayaan Perempuan dan Perlindungan Anak",
    "Bidang Pangan", "Bidang Pertanahan", "Bidang Lingkungan Hidup",
    "Bidang Administrasi Kependudukan dan Pencatatan Sipil", "Bidang Pemberdayaan Masyarakat dan Desa",
    "Bidang Pengendalian Penduduk dan Keluarga Berencana", "Bidang Perhubungan",
    "Bidang Komunikasi dan Informatika", "Bidang Koperasi, Usaha Kecil dan Menengah",
    "Bidang Penanaman Modal", "Bidang Kepemudaan dan Olahraga", "Bidang Statistik",
    "Bidang Persandian", "Bidang Kebudayaan", "Bidang Perpustakaan", "Bidang Kearsipan",
    "Bidang Kelautan dan Perikanan", "Bidang Pariwisata", "Bidang Pertanian", "Bidang Kehutanan",
    "Bidang Energi dan Sumber Daya Mineral", "Bidang Perdagangan", "Bidang Perindustrian", "Bidang Transmigrasi",
]
TRANTIB = "Bidang Ketentraman dan Ketertiban Umum serta Perlindungan Masyarakat"

def _norm(s):
    s = re.sub(r'\bbidang\b', '', str(s).lower())
    s = re.sub(r'[^a-z0-9 ]', ' ', s)
    return re.sub(r'\s+', ' ', s).strip()

URUSAN_CANON = {_norm(u): u for u in URUSAN_STATIC}

def canonical_urusan(name: str):
    n = _norm(name)
    if "kebakaran" in n:
        return TRANTIB, "Sub Urusan Kebakaran"
    if "ketenteraman" in n or "ketertiban" in n:
        return TRANTIB, "Urusan Ketentraman dan Ketertiban Umum"
    if n in URUSAN_CANON:
        return URUSAN_CANON[n], None
    return name.strip(), None

def parse_indicator_items(text, n):
    text = re.sub(r'\s+', ' ', str(text)).strip()
    marks = []
    idx = 0
    for k in range(1, int(n) + 1):
        m = re.compile(r'\b' + str(k) + r'[\.\)]?\s+').search(text, idx)
        if not m:
            break
        marks.append((m.start(), m.end()))
        idx = m.end()
    items = []
    for i, (st, en) in enumerate(marks):
        nxt = marks[i + 1][0] if i + 1 < len(marks) else len(text)
        seg = text[en:nxt].strip().strip(',;.').strip()
        if seg:
            items.append(seg)
    return items

async def import_indicators_from_workbook(wb):
    sheet_level = {"Provinsi": "provinsi", "Kabupaten Kota": "kabupaten", "Kabupaten/Kota": "kabupaten"}
    inserted = 0
    for sn in wb.sheetnames:
        level = sheet_level.get(sn)
        if not level:
            continue
        ws = wb[sn]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 4:
                continue
            kategori, name, count, daftar = row[0], row[1], row[2], row[3]
            if not name or not daftar or not str(kategori or "").lower().startswith("faktor teknis"):
                continue
            try:
                n = int(count or 0)
            except Exception:
                n = 0
            items = parse_indicator_items(daftar, n) if n else []
            if not items:
                continue
            urusan, sub = canonical_urusan(name)
            await db.indicators.delete_many({"type": "teknis", "level": level, "urusan": urusan, "sub_urusan": sub})
            for order, item in enumerate(items, start=1):
                await db.indicators.insert_one({"id": str(uuid.uuid4()), "type": "teknis", "name": item,
                    "description": "", "level": level, "urusan": urusan, "sub_urusan": sub,
                    "order": order, "created_at": now_iso()})
                inserted += 1
    return inserted


def tipe_from_score(v: float) -> dict:
    if v > 800: return {"key": "A", "label": "Tipe A"}
    if v > 600: return {"key": "B", "label": "Tipe B"}
    if v > 400: return {"key": "C", "label": "Tipe C"}
    if v > 300: return {"key": "BIDANG", "label": "Setingkat Bidang"}
    return {"key": "SUBBIDANG", "label": "Setingkat Subbidang/Seksi"}

now_iso = lambda: datetime.now(timezone.utc).isoformat()

app = FastAPI()
api_router = APIRouter(prefix="/api")

# ---------------- Models ----------------
class LoginInput(BaseModel):
    email: EmailStr
    password: str

class UserCreate(BaseModel):
    email: EmailStr
    password: str
    name: str
    role: str
    area: Optional[str] = None

class UserUpdate(BaseModel):
    name: Optional[str] = None
    role: Optional[str] = None
    area: Optional[str] = None
    active: Optional[bool] = None

class PasswordReset(BaseModel):
    password: str

class IndicatorInput(BaseModel):
    type: str  # umum | teknis
    name: str
    description: Optional[str] = ""
    level: Optional[str] = None   # provinsi | kabupaten (teknis)
    urusan: Optional[str] = None
    sub_urusan: Optional[str] = None
    order: Optional[int] = 0

class PeriodInput(BaseModel):
    year: int
    name: str
    start_date: str
    end_date: str
    upload_locked: bool = False
    active: bool = True

class SubmissionCreate(BaseModel):
    device_name: str
    urusan: str
    sub_urusan: Optional[str] = None
    period_id: str

class VerifyInput(BaseModel):
    action: str
    notes: Optional[str] = ""

class ScoreItem(BaseModel):
    indicator_id: str
    ok: bool = True
    note: Optional[str] = ""
    data_validasi: Optional[str] = ""
    score: float

class ScoreInput(BaseModel):
    items: List[ScoreItem]
    overall_note: Optional[str] = ""

class ReportInput(BaseModel):
    area: str
    device_name: str
    apply_multiplier: bool = False

# ---------------- Auth routes ----------------
@api_router.post("/auth/login")
async def login(input: LoginInput, response: Response):
    email = input.email.lower()
    user = await db.users.find_one({"email": email})
    if not user or not verify_password(input.password, user.get("password_hash", "")):
        raise HTTPException(status_code=401, detail="Email atau kata sandi salah")
    if not user.get("active", True):
        raise HTTPException(status_code=403, detail="Akun dinonaktifkan")
    token = create_access_token(user["id"])
    set_auth_cookie(response, token)
    return {"id": user["id"], "email": user["email"], "name": user["name"], "role": user["role"],
            "area": user.get("area"), "level": AREA_LEVEL.get(user.get("area")), "token": token}

@api_router.post("/auth/logout")
async def logout(response: Response, user: dict = Depends(get_current_user)):
    response.delete_cookie("access_token", path="/")
    return {"message": "Berhasil keluar"}

@api_router.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    user["level"] = AREA_LEVEL.get(user.get("area"))
    return user

@api_router.get("/reference/areas")
async def ref_areas(user: dict = Depends(get_current_user)):
    return AREAS

@api_router.get("/reference/urusan")
async def ref_urusan(user: dict = Depends(get_current_user)):
    teknis = await db.indicators.find({"type": "teknis"}, {"_id": 0, "urusan": 1, "sub_urusan": 1}).to_list(5000)
    present = {}
    for t in teknis:
        present.setdefault(t["urusan"], set())
        if t.get("sub_urusan"):
            present[t["urusan"]].add(t["sub_urusan"])
    # order: static list first (those present), then extras sorted
    ordered = [u for u in URUSAN_STATIC if u in present]
    extras = sorted([u for u in present.keys() if u not in URUSAN_STATIC])
    urusan_list = ordered + extras
    if not urusan_list:
        urusan_list = URUSAN_STATIC
    sub_map = {u: sorted(list(s)) for u, s in present.items() if s}
    return {"urusan": urusan_list, "sub_urusan": sub_map}

@api_router.post("/indicators/import")
async def import_indicators(file: UploadFile = File(...), user: dict = Depends(require_roles("admin"))):
    from openpyxl import load_workbook
    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Unggah berkas Excel (.xlsx)")
    data = await file.read()
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Gagal membaca Excel: {e}")
    inserted = await import_indicators_from_workbook(wb)
    ref = await ref_urusan(user)
    return {"message": f"Impor selesai: {inserted} indikator teknis dimuat", "inserted": inserted, "urusan_count": len(ref["urusan"])}

# ---------------- Users ----------------
@api_router.get("/users")
async def list_users(user: dict = Depends(require_roles("admin"))):
    return await db.users.find({}, {"_id": 0, "password_hash": 0}).to_list(1000)

@api_router.post("/users")
async def create_user(input: UserCreate, user: dict = Depends(require_roles("admin"))):
    email = input.email.lower()
    if input.role not in ("admin", "perangkat", "verifikator", "penilai"):
        raise HTTPException(status_code=400, detail="Peran tidak valid")
    if input.role in ("perangkat", "verifikator") and not input.area:
        raise HTTPException(status_code=400, detail="Area wajib untuk peran perangkat/verifikator")
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Email sudah terdaftar")
    doc = {"id": str(uuid.uuid4()), "email": email, "name": input.name, "role": input.role,
           "area": input.area, "active": True, "password_hash": hash_password(input.password), "created_at": now_iso()}
    await db.users.insert_one(doc)
    return {k: v for k, v in doc.items() if k not in ("_id", "password_hash")}

@api_router.put("/users/{uid}")
async def update_user(uid: str, input: UserUpdate, user: dict = Depends(require_roles("admin"))):
    updates = {k: v for k, v in input.model_dump().items() if v is not None}
    if not updates: raise HTTPException(status_code=400, detail="Tidak ada perubahan")
    res = await db.users.update_one({"id": uid}, {"$set": updates})
    if res.matched_count == 0: raise HTTPException(status_code=404, detail="Pengguna tidak ditemukan")
    return await db.users.find_one({"id": uid}, {"_id": 0, "password_hash": 0})

@api_router.post("/users/{uid}/reset-password")
async def reset_pw(uid: str, input: PasswordReset, user: dict = Depends(require_roles("admin"))):
    res = await db.users.update_one({"id": uid}, {"$set": {"password_hash": hash_password(input.password)}})
    if res.matched_count == 0: raise HTTPException(status_code=404, detail="Pengguna tidak ditemukan")
    return {"message": "Kata sandi diperbarui"}

@api_router.delete("/users/{uid}")
async def delete_user(uid: str, user: dict = Depends(require_roles("admin"))):
    if uid == user["id"]: raise HTTPException(status_code=400, detail="Tidak dapat menghapus diri sendiri")
    await db.users.delete_one({"id": uid})
    return {"message": "Pengguna dihapus"}

# ---------------- Indicators ----------------
@api_router.get("/indicators")
async def list_indicators(type: Optional[str] = None, level: Optional[str] = None,
                          urusan: Optional[str] = None, user: dict = Depends(get_current_user)):
    q = {}
    if type: q["type"] = type
    if level: q["level"] = level
    if urusan: q["urusan"] = urusan
    return await db.indicators.find(q, {"_id": 0}).sort("order", 1).to_list(2000)

@api_router.get("/indicators/for-submission")
async def indicators_for_submission(level: str, urusan: str, sub_urusan: Optional[str] = None,
                                    user: dict = Depends(get_current_user)):
    umum = await db.indicators.find({"type": "umum"}, {"_id": 0}).sort("order", 1).to_list(100)
    tq = {"type": "teknis", "level": level, "urusan": urusan}
    if sub_urusan:
        tq["sub_urusan"] = sub_urusan
    else:
        tq["sub_urusan"] = None
    teknis = await db.indicators.find(tq, {"_id": 0}).sort("order", 1).to_list(500)
    return {"umum": umum, "teknis": teknis}

@api_router.post("/indicators")
async def create_indicator(input: IndicatorInput, user: dict = Depends(require_roles("admin"))):
    if input.type not in ("umum", "teknis"):
        raise HTTPException(status_code=400, detail="Tipe indikator tidak valid")
    count = await db.indicators.count_documents({"type": input.type, "urusan": input.urusan, "level": input.level})
    doc = {"id": str(uuid.uuid4()), **input.model_dump(), "created_at": now_iso()}
    if not doc.get("order"): doc["order"] = count + 1
    await db.indicators.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}

@api_router.put("/indicators/{iid}")
async def update_indicator(iid: str, input: IndicatorInput, user: dict = Depends(require_roles("admin"))):
    res = await db.indicators.update_one({"id": iid}, {"$set": input.model_dump()})
    if res.matched_count == 0: raise HTTPException(status_code=404, detail="Indikator tidak ditemukan")
    return await db.indicators.find_one({"id": iid}, {"_id": 0})

@api_router.delete("/indicators/{iid}")
async def delete_indicator(iid: str, user: dict = Depends(require_roles("admin"))):
    await db.indicators.delete_one({"id": iid})
    return {"message": "Indikator dihapus"}

# ---------------- Periods & Notifications ----------------
@api_router.get("/periods")
async def list_periods(user: dict = Depends(get_current_user)):
    return await db.periods.find({}, {"_id": 0}).sort("year", -1).to_list(1000)

@api_router.get("/periods/active")
async def active_period(user: dict = Depends(get_current_user)):
    return await db.periods.find_one({"active": True}, {"_id": 0}) or {}

@api_router.post("/periods")
async def create_period(input: PeriodInput, user: dict = Depends(require_roles("admin"))):
    if input.active: await db.periods.update_many({}, {"$set": {"active": False}})
    doc = {"id": str(uuid.uuid4()), **input.model_dump(), "created_at": now_iso()}
    await db.periods.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}

@api_router.put("/periods/{pid}")
async def update_period(pid: str, input: PeriodInput, user: dict = Depends(require_roles("admin"))):
    if input.active: await db.periods.update_many({"id": {"$ne": pid}}, {"$set": {"active": False}})
    res = await db.periods.update_one({"id": pid}, {"$set": input.model_dump()})
    if res.matched_count == 0: raise HTTPException(status_code=404, detail="Periode tidak ditemukan")
    return await db.periods.find_one({"id": pid}, {"_id": 0})

@api_router.delete("/periods/{pid}")
async def delete_period(pid: str, user: dict = Depends(require_roles("admin"))):
    await db.periods.delete_one({"id": pid})
    return {"message": "Periode dihapus"}

@api_router.get("/notifications")
async def notifications(user: dict = Depends(get_current_user)):
    p = await db.periods.find_one({"active": True}, {"_id": 0})
    out = []
    if not p:
        return [{"level": "info", "message": "Belum ada periode evaluasi aktif."}]
    today = datetime.now(timezone.utc).date()
    try:
        start = datetime.strptime(p["start_date"], "%Y-%m-%d").date()
        end = datetime.strptime(p["end_date"], "%Y-%m-%d").date()
    except Exception:
        return out
    if today < start:
        d = (start - today).days
        if d <= 14:
            out.append({"level": "info", "message": f"Evaluasi '{p['name']}' akan dibuka dalam {d} hari ({p['start_date']})."})
    elif start <= today <= end:
        d = (end - today).days
        if p.get("upload_locked"):
            out.append({"level": "warning", "message": "Unggahan berkas sedang DIKUNCI oleh administrator."})
        if d <= 7:
            out.append({"level": "urgent", "message": f"Batas akhir evaluasi {d} hari lagi ({p['end_date']}). Segera lengkapi & kirim berkas!"})
        else:
            out.append({"level": "info", "message": f"Periode evaluasi '{p['name']}' sedang berlangsung. Batas {p['end_date']} ({d} hari lagi)."})
    else:
        out.append({"level": "warning", "message": f"Periode evaluasi '{p['name']}' telah berakhir ({p['end_date']})."})
    return out

# ---------------- Submissions ----------------
def scope_query(user):
    role = user["role"]
    if role == "perangkat":
        return {"perangkat_user_id": user["id"]}
    if role == "verifikator":
        return {"area": user.get("area"), "status": {"$in": ["menunggu_verifikasi", "ditolak", "menunggu_penilaian", "selesai"]}}
    if role == "penilai":
        return {"status": {"$in": ["menunggu_penilaian", "selesai"]}}
    return {}

async def add_audit(sid, action, user, detail=""):
    entry = {"action": action, "by": user["name"], "by_role": user["role"], "at": now_iso(), "detail": detail}
    await db.submissions.update_one({"id": sid}, {"$push": {"audit": entry}})

@api_router.get("/submissions")
async def list_submissions(user: dict = Depends(get_current_user)):
    return await db.submissions.find(scope_query(user), {"_id": 0}).sort("created_at", -1).to_list(3000)

@api_router.get("/submissions/{sid}")
async def get_submission(sid: str, user: dict = Depends(get_current_user)):
    s = await db.submissions.find_one({"id": sid}, {"_id": 0})
    if not s: raise HTTPException(status_code=404, detail="Pengajuan tidak ditemukan")
    if user["role"] == "perangkat" and s["perangkat_user_id"] != user["id"]:
        raise HTTPException(status_code=403, detail="Akses ditolak")
    if user["role"] == "verifikator" and s["area"] != user.get("area"):
        raise HTTPException(status_code=403, detail="Akses ditolak (area berbeda)")
    return s

@api_router.post("/submissions")
async def create_submission(input: SubmissionCreate, user: dict = Depends(require_roles("perangkat"))):
    if not user.get("area"):
        raise HTTPException(status_code=400, detail="Akun Anda belum memiliki area")
    period = await db.periods.find_one({"id": input.period_id})
    if not period: raise HTTPException(status_code=404, detail="Periode tidak ditemukan")
    level = AREA_LEVEL.get(user["area"])
    # copy faktor umum from latest submission in same period (upload once)
    prev = await db.submissions.find_one(
        {"perangkat_user_id": user["id"], "period_id": period["id"]},
        {"_id": 0}, sort=[("created_at", -1)])
    umum_ids = {i["id"] for i in await db.indicators.find({"type": "umum"}).to_list(100)}
    uploads = {}
    if prev:
        for k, v in (prev.get("uploads") or {}).items():
            if k in umum_ids: uploads[k] = v
    doc = {"id": str(uuid.uuid4()), "period_id": period["id"], "year": period["year"],
           "area": user["area"], "level": level, "device_name": input.device_name,
           "urusan": input.urusan, "sub_urusan": input.sub_urusan,
           "perangkat_user_id": user["id"], "perangkat_name": user["name"],
           "status": "draft", "uploads": uploads, "verification": None, "scoring": None,
           "rejection_note": None, "history": [{"status": "draft", "at": now_iso(), "by": user["name"]}],
           "audit": [{"action": "Membuat pengajuan", "by": user["name"], "by_role": "perangkat", "at": now_iso(), "detail": f"{input.urusan}"}],
           "created_at": now_iso(), "updated_at": now_iso()}
    await db.submissions.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}

async def check_upload_allowed(period_id: str):
    period = await db.periods.find_one({"id": period_id})
    if not period: raise HTTPException(status_code=404, detail="Periode tidak ditemukan")
    if period.get("upload_locked"): raise HTTPException(status_code=403, detail="Unggahan berkas telah dikunci oleh admin")
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    if period.get("end_date") and today > period["end_date"]:
        raise HTTPException(status_code=403, detail="Periode unggah berkas telah berakhir")

async def required_indicator_ids(s):
    umum = await db.indicators.find({"type": "umum"}).to_list(100)
    tq = {"type": "teknis", "level": s["level"], "urusan": s["urusan"], "sub_urusan": s.get("sub_urusan")}
    teknis = await db.indicators.find(tq).to_list(500)
    return [i["id"] for i in umum] + [i["id"] for i in teknis], umum, teknis

@api_router.post("/submissions/{sid}/upload")
async def upload_evidence(sid: str, indicator_id: str = Form(...), file: UploadFile = File(...),
                          user: dict = Depends(require_roles("perangkat"))):
    s = await db.submissions.find_one({"id": sid})
    if not s or s["perangkat_user_id"] != user["id"]:
        raise HTTPException(status_code=404, detail="Pengajuan tidak ditemukan")
    if s["status"] not in ("draft", "ditolak"):
        raise HTTPException(status_code=403, detail="Pengajuan tidak dapat diubah pada status ini")
    await check_upload_allowed(s["period_id"])
    ext = file.filename.split(".")[-1].lower() if "." in file.filename else "bin"
    if ext not in ("pdf", "png", "jpg", "jpeg", "webp", "doc", "docx", "xls", "xlsx"):
        raise HTTPException(status_code=400, detail="Format berkas tidak didukung")
    data = await file.read()
    if len(data) > 15 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Ukuran berkas maksimal 15MB")
    file_id = str(uuid.uuid4())
    path = f"{APP_NAME}/{user['id']}/{sid}/{file_id}.{ext}"
    result = put_object(path, data, file.content_type or "application/octet-stream")
    await db.files.insert_one({"id": file_id, "storage_path": result["path"], "original_filename": file.filename,
                               "content_type": file.content_type, "submission_id": sid, "is_deleted": False, "created_at": now_iso()})
    meta = {"file_id": file_id, "original_filename": file.filename, "uploaded_at": now_iso()}
    await db.submissions.update_one({"id": sid}, {"$set": {f"uploads.{indicator_id}": meta, "updated_at": now_iso()}})
    return meta

@api_router.get("/files/{file_id}/download")
async def download_file(file_id: str, user: dict = Depends(get_current_user)):
    from fastapi import Response as FResponse
    record = await db.files.find_one({"id": file_id, "is_deleted": False})
    if not record: raise HTTPException(status_code=404, detail="Berkas tidak ditemukan")
    data, ct = get_object(record["storage_path"])
    return FResponse(content=data, media_type=record.get("content_type") or ct,
                     headers={"Content-Disposition": f'inline; filename="{record["original_filename"]}"'})

@api_router.post("/submissions/{sid}/submit")
async def submit_submission(sid: str, user: dict = Depends(require_roles("perangkat"))):
    s = await db.submissions.find_one({"id": sid})
    if not s or s["perangkat_user_id"] != user["id"]:
        raise HTTPException(status_code=404, detail="Pengajuan tidak ditemukan")
    if s["status"] not in ("draft", "ditolak"):
        raise HTTPException(status_code=403, detail="Pengajuan sudah dikirim")
    await check_upload_allowed(s["period_id"])
    req_ids, umum, teknis = await required_indicator_ids(s)
    id_name = {i["id"]: i["name"] for i in umum + teknis}
    missing = [id_name.get(i, i) for i in req_ids if i not in (s.get("uploads") or {})]
    if missing:
        raise HTTPException(status_code=400, detail=f"Berkas belum lengkap: {', '.join(missing)}")
    hist = s.get("history", []); hist.append({"status": "menunggu_verifikasi", "at": now_iso(), "by": user["name"]})
    await db.submissions.update_one({"id": sid}, {"$set": {"status": "menunggu_verifikasi", "rejection_note": None, "updated_at": now_iso(), "history": hist}})
    await add_audit(sid, "Mengirim untuk verifikasi", user)
    return {"message": "Pengajuan dikirim untuk verifikasi"}

@api_router.post("/submissions/{sid}/verify")
async def verify_submission(sid: str, input: VerifyInput, user: dict = Depends(require_roles("verifikator"))):
    s = await db.submissions.find_one({"id": sid})
    if not s: raise HTTPException(status_code=404, detail="Pengajuan tidak ditemukan")
    if s["area"] != user.get("area"): raise HTTPException(status_code=403, detail="Akses ditolak (area berbeda)")
    if s["status"] != "menunggu_verifikasi": raise HTTPException(status_code=403, detail="Pengajuan tidak dalam status verifikasi")
    hist = s.get("history", [])
    if input.action == "approve":
        hist.append({"status": "menunggu_penilaian", "at": now_iso(), "by": user["name"]})
        await db.submissions.update_one({"id": sid}, {"$set": {"status": "menunggu_penilaian",
            "verification": {"verifikator_id": user["id"], "verifikator_name": user["name"], "notes": input.notes, "verified_at": now_iso()},
            "updated_at": now_iso(), "history": hist}})
        await add_audit(sid, "Verifikasi disetujui", user, input.notes or "")
        return {"message": "Terverifikasi, diteruskan ke penilai"}
    elif input.action == "reject":
        if not input.notes: raise HTTPException(status_code=400, detail="Catatan perbaikan wajib diisi saat menolak")
        hist.append({"status": "ditolak", "at": now_iso(), "by": user["name"], "note": input.notes})
        await db.submissions.update_one({"id": sid}, {"$set": {"status": "ditolak", "rejection_note": input.notes,
            "verification": {"verifikator_id": user["id"], "verifikator_name": user["name"], "notes": input.notes, "verified_at": now_iso()},
            "updated_at": now_iso(), "history": hist}})
        await add_audit(sid, "Dikembalikan untuk perbaikan", user, input.notes)
        return {"message": "Dikembalikan ke perangkat untuk perbaikan"}
    raise HTTPException(status_code=400, detail="Aksi tidak valid")

@api_router.post("/submissions/{sid}/score")
async def score_submission(sid: str, input: ScoreInput, user: dict = Depends(require_roles("penilai"))):
    s = await db.submissions.find_one({"id": sid})
    if not s: raise HTTPException(status_code=404, detail="Pengajuan tidak ditemukan")
    if s["status"] != "menunggu_penilaian": raise HTTPException(status_code=403, detail="Pengajuan belum terverifikasi")
    req_ids, umum, teknis = await required_indicator_ids(s)
    umum_ids = {i["id"] for i in umum}; teknis_ids = {i["id"] for i in teknis}
    name_map = {i["id"]: i["name"] for i in umum + teknis}
    submitted = {it.indicator_id: it for it in input.items}
    if set(submitted.keys()) != set(req_ids):
        raise HTTPException(status_code=400, detail="Semua indikator harus dinilai tepat satu kali")
    validations = {}
    umum_scores, teknis_scores = [], []
    for iid, it in submitted.items():
        if it.score < 0 or it.score > 1000:
            raise HTTPException(status_code=400, detail="Skor harus 0 - 1000")
        validations[iid] = {"indicator_name": name_map.get(iid, ""), "ok": it.ok, "note": it.note,
                            "data_validasi": it.data_validasi, "score": it.score}
        (umum_scores if iid in umum_ids else teknis_scores).append(it.score)
    umum_avg = round(sum(umum_scores) / len(umum_scores), 2) if umum_scores else 0
    teknis_avg = round(sum(teknis_scores) / len(teknis_scores), 2) if teknis_scores else 0
    scoring = {"penilai_id": user["id"], "penilai_name": user["name"], "validations": validations,
               "umum_avg": umum_avg, "teknis_avg": teknis_avg, "overall_note": input.overall_note, "scored_at": now_iso()}
    hist = s.get("history", []); hist.append({"status": "selesai", "at": now_iso(), "by": user["name"]})
    await db.submissions.update_one({"id": sid}, {"$set": {"status": "selesai", "scoring": scoring, "updated_at": now_iso(), "history": hist}})
    await add_audit(sid, "Penilaian selesai", user, f"Umum {umum_avg} / Teknis {teknis_avg}")
    return {"message": "Penilaian selesai", "scoring": scoring}

# ---------------- Reports ----------------
async def build_report_rows(area, device_name):
    subs = await db.submissions.find({"area": area, "device_name": device_name, "status": "selesai"}, {"_id": 0}).to_list(500)
    rows = []
    for s in subs:
        sc = s.get("scoring") or {}
        umum = sc.get("umum_avg", 0); teknis = sc.get("teknis_avg", 0)
        total = round(0.2 * umum + 0.8 * teknis, 2)
        rows.append({"submission_id": s["id"], "urusan": s["urusan"], "sub_urusan": s.get("sub_urusan"),
                     "umum_avg": umum, "teknis_avg": teknis, "total": total})
    return rows

@api_router.get("/reports/perangkat-list")
async def report_perangkat_list(user: dict = Depends(require_roles("penilai", "admin"))):
    subs = await db.submissions.find({"status": "selesai"}, {"_id": 0, "area": 1, "device_name": 1}).to_list(3000)
    seen = {}
    for s in subs:
        seen[(s["area"], s["device_name"])] = True
    return [{"area": a, "device_name": d} for (a, d) in seen.keys()]

@api_router.get("/reports/preview")
async def report_preview(area: str, device_name: str, apply_multiplier: bool = False,
                         user: dict = Depends(require_roles("penilai", "admin"))):
    rows = await build_report_rows(area, device_name)
    for r in rows:
        final = round(r["total"] * 1.1, 2) if apply_multiplier else r["total"]
        r["final"] = final
        r["tipe"] = tipe_from_score(final)
    return {"area": area, "device_name": device_name, "apply_multiplier": apply_multiplier, "rows": rows}

@api_router.post("/reports")
async def create_report(input: ReportInput, user: dict = Depends(require_roles("penilai", "admin"))):
    rows = await build_report_rows(input.area, input.device_name)
    if not rows: raise HTTPException(status_code=400, detail="Tidak ada pengajuan selesai untuk perangkat ini")
    for r in rows:
        final = round(r["total"] * 1.1, 2) if input.apply_multiplier else r["total"]
        r["final"] = final; r["tipe"] = tipe_from_score(final)
    doc = {"id": str(uuid.uuid4()), "area": input.area, "device_name": input.device_name,
           "level": AREA_LEVEL.get(input.area), "apply_multiplier": input.apply_multiplier, "rows": rows,
           "penilai_id": user["id"], "penilai_name": user["name"], "created_at": now_iso()}
    await db.reports.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}

@api_router.get("/reports")
async def list_reports(user: dict = Depends(get_current_user)):
    q = {}
    if user["role"] == "perangkat":
        q["area"] = user.get("area")
    elif user["role"] == "verifikator":
        q["area"] = user.get("area")
    return await db.reports.find(q, {"_id": 0}).sort("created_at", -1).to_list(1000)

@api_router.delete("/reports/{rid}")
async def delete_report(rid: str, user: dict = Depends(require_roles("penilai", "admin"))):
    await db.reports.delete_one({"id": rid})
    return {"message": "Laporan dihapus"}

@api_router.get("/reports/{rid}/excel")
async def report_excel(rid: str, auth: str = Query(None), authorization: str = Header(None)):
    token = auth or (authorization[7:] if authorization and authorization.startswith("Bearer ") else None)
    if not token: raise HTTPException(status_code=401, detail="Tidak terautentikasi")
    try:
        payload = jwt.decode(token, get_jwt_secret(), algorithms=[JWT_ALGORITHM])
        u = await db.users.find_one({"id": payload["sub"]})
        if not u: raise HTTPException(status_code=401, detail="Pengguna tidak ditemukan")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")
    r = await db.reports.find_one({"id": rid}, {"_id": 0})
    if not r: raise HTTPException(status_code=404, detail="Laporan tidak ditemukan")
    if u["role"] in ("perangkat", "verifikator") and r["area"] != u.get("area"):
        raise HTTPException(status_code=403, detail="Akses ditolak")
    wb = Workbook(); ws = wb.active; ws.title = "Hasil Penilaian"
    head_fill = PatternFill("solid", fgColor="064E3B"); head_font = Font(bold=True, color="FFFFFF")
    ws.append([f"Laporan Hasil Penilaian PP 18/2016"])
    ws.append([f"Area: {r['area']} ({r.get('level','')})"])
    ws.append([f"Perangkat Daerah: {r['device_name']}"])
    ws.append([f"Pengali 1,1: {'Ya' if r['apply_multiplier'] else 'Tidak'}"])
    ws.append([])
    headers = ["No", "Urusan", "Sub Urusan", "Nilai Faktor Umum (20%)", "Nilai Faktor Teknis (80%)", "Total", "Nilai Akhir", "Tipe"]
    ws.append(headers)
    hr = ws.max_row
    for c in range(1, len(headers) + 1):
        ws.cell(row=hr, column=c).fill = head_fill
        ws.cell(row=hr, column=c).font = head_font
        ws.cell(row=hr, column=c).alignment = Alignment(horizontal="center", wrap_text=True)
    for idx, row in enumerate(r["rows"], start=1):
        ws.append([idx, row["urusan"], row.get("sub_urusan") or "-", row["umum_avg"], row["teknis_avg"],
                   row["total"], row["final"], row["tipe"]["label"]])
    widths = [5, 40, 28, 20, 20, 12, 12, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    fname = f"Laporan_{r['device_name'].replace(' ', '_')}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})

# ---------------- Stats ----------------
@api_router.get("/stats/overview")
async def stats_overview(user: dict = Depends(get_current_user)):
    base = scope_query(user)
    base = {k: v for k, v in base.items() if k != "status"}
    counts = {}
    for st in ["draft", "menunggu_verifikasi", "ditolak", "menunggu_penilaian", "selesai"]:
        counts[st] = await db.submissions.count_documents({**base, "status": st})
    counts["total"] = await db.submissions.count_documents(base)
    extra = {}
    if user["role"] == "admin":
        extra["users"] = await db.users.count_documents({})
        extra["indicators"] = await db.indicators.count_documents({})
        extra["reports"] = await db.reports.count_documents({})
    return {"counts": counts, **extra}

@api_router.get("/stats/yearly")
async def stats_yearly(user: dict = Depends(get_current_user)):
    q = {"status": "selesai"}
    if user["role"] == "perangkat": q["perangkat_user_id"] = user["id"]
    elif user["role"] == "verifikator": q["area"] = user.get("area")
    subs = await db.submissions.find(q, {"_id": 0}).to_list(5000)
    by_year = {}
    for s in subs:
        sc = s.get("scoring") or {}
        total = round(0.2 * sc.get("umum_avg", 0) + 0.8 * sc.get("teknis_avg", 0), 2)
        by_year.setdefault(s["year"], []).append(total)
    return [{"year": y, "avg_total": round(sum(v) / len(v), 2), "count": len(v)} for y, v in sorted(by_year.items())]

@api_router.get("/stats/by-area")
async def stats_by_area(user: dict = Depends(require_roles("admin", "penilai"))):
    subs = await db.submissions.find({"status": "selesai"}, {"_id": 0}).to_list(5000)
    by_area = {}
    for s in subs:
        sc = s.get("scoring") or {}
        total = round(0.2 * sc.get("umum_avg", 0) + 0.8 * sc.get("teknis_avg", 0), 2)
        by_area.setdefault(s["area"], []).append(total)
    return [{"area": a, "avg_total": round(sum(v) / len(v), 2), "count": len(v)} for a, v in by_area.items()]

@api_router.get("/audit")
async def audit_feed(user: dict = Depends(require_roles("admin"))):
    subs = await db.submissions.find({}, {"_id": 0, "id": 1, "device_name": 1, "area": 1, "urusan": 1, "audit": 1}).to_list(3000)
    feed = []
    for s in subs:
        for a in (s.get("audit") or []):
            feed.append({**a, "device_name": s["device_name"], "area": s["area"], "urusan": s["urusan"], "submission_id": s["id"]})
    feed.sort(key=lambda x: x.get("at", ""), reverse=True)
    return feed[:200]

# ---------------- Seed ----------------
UMUM = [("Jumlah Penduduk", "Total penduduk daerah (jiwa) - karakteristik daerah."),
        ("Luas Wilayah", "Luas wilayah administratif daerah (km2)."),
        ("Jumlah APBD", "Besaran Anggaran Pendapatan dan Belanja Daerah (Rp).")]

TEKNIS_SEED = {
    ("kabupaten", "Bidang Pendidikan", None): [
        "Jumlah satuan pendidikan (SD/SMP dan sederajat)",
        "Jumlah anak usia sekolah (usia 7-15 tahun)",
        "Jumlah Kurikulum muatan lokal yang dikembangkan",
    ],
    ("provinsi", "Bidang Pendidikan", None): [
        "Jumlah satuan pendidikan menengah (SMA/SMK dan sederajat)",
        "Jumlah anak usia sekolah menengah (usia 16-18 tahun)",
        "Jumlah Kurikulum muatan lokal pendidikan menengah",
    ],
    ("kabupaten", "Bidang Kesehatan", None): [
        "Jumlah penduduk yang dilayani fasilitas kesehatan",
        "Jumlah fasilitas kesehatan (Puskesmas/RSUD)",
        "Cakupan pelayanan kesehatan dasar",
    ],
    ("provinsi", "Bidang Kesehatan", None): [
        "Jumlah penduduk yang dilayani rujukan provinsi",
        "Jumlah rumah sakit rujukan kewenangan provinsi",
        "Cakupan pelayanan kesehatan rujukan",
    ],
    ("kabupaten", "Bidang Pekerjaan Umum dan Penataan Ruang", None): [
        "Panjang jalan kewenangan kabupaten/kota",
        "Jumlah bangunan gedung strategis daerah",
        "Luas kawasan permukiman yang ditangani",
    ],
    ("provinsi", "Bidang Pekerjaan Umum dan Penataan Ruang", None): [
        "Panjang jalan kewenangan provinsi",
        "Panjang sungai lintas kabupaten/kota",
        "Luas kawasan strategis provinsi yang ditangani",
    ],
    ("kabupaten", "Bidang Ketentraman dan Ketertiban Umum serta Perlindungan Masyarakat", "Urusan Ketentraman dan Ketertiban Umum"): [
        "Jumlah pelanggaran Peraturan Daerah",
        "Jumlah personel Satuan Polisi Pamong Praja",
        "Cakupan penegakan Perda dan Perkada",
    ],
    ("kabupaten", "Bidang Ketentraman dan Ketertiban Umum serta Perlindungan Masyarakat", "Sub Urusan Kebakaran"): [
        "Jumlah kejadian kebakaran per tahun",
        "Luas wilayah layanan pemadam kebakaran",
        "Jumlah armada dan personel pemadam kebakaran",
    ],
}

async def seed():
    await db.users.create_index("email", unique=True)
    await db.users.create_index("id", unique=True)
    admin_email = os.environ.get("ADMIN_EMAIL", "admin@example.com").lower()
    admin_pw = os.environ.get("ADMIN_PASSWORD", "admin123")
    existing = await db.users.find_one({"email": admin_email})
    if not existing:
        await db.users.insert_one({"id": str(uuid.uuid4()), "email": admin_email, "name": "Administrator", "role": "admin",
                                   "area": None, "active": True, "password_hash": hash_password(admin_pw), "created_at": now_iso()})
    elif not verify_password(admin_pw, existing.get("password_hash", "")):
        await db.users.update_one({"email": admin_email}, {"$set": {"password_hash": hash_password(admin_pw)}})
    demo = [
        ("verifikator@kalteng.go.id", "Verif123!", "Verifikator Palangka Raya", "verifikator", "Kota Palangka Raya"),
        ("penilai@kalteng.go.id", "Nilai123!", "Tim Penilai Provinsi", "penilai", None),
        ("perangkat@kalteng.go.id", "Kerja123!", "Dinas Pendidikan Kota Palangka Raya", "perangkat", "Kota Palangka Raya"),
    ]
    for email, pw, name, role, area in demo:
        if not await db.users.find_one({"email": email}):
            await db.users.insert_one({"id": str(uuid.uuid4()), "email": email, "name": name, "role": role,
                                       "area": area, "active": True, "password_hash": hash_password(pw), "created_at": now_iso()})
    # indicators
    if await db.indicators.count_documents({"type": "umum"}) == 0:
        for idx, (n, d) in enumerate(UMUM, start=1):
            await db.indicators.insert_one({"id": str(uuid.uuid4()), "type": "umum", "name": n, "description": d,
                                            "level": None, "urusan": None, "sub_urusan": None, "order": idx, "created_at": now_iso()})
    if await db.indicators.count_documents({"type": "teknis"}) == 0:
        for (level, urusan, sub), names in TEKNIS_SEED.items():
            for idx, n in enumerate(names, start=1):
                await db.indicators.insert_one({"id": str(uuid.uuid4()), "type": "teknis", "name": n, "description": "",
                                                "level": level, "urusan": urusan, "sub_urusan": sub, "order": idx, "created_at": now_iso()})
    if await db.periods.count_documents({}) == 0:
        y = datetime.now(timezone.utc).year
        await db.periods.insert_one({"id": str(uuid.uuid4()), "year": y, "name": f"Evaluasi Perangkat Daerah {y}",
                                     "start_date": f"{y}-01-01", "end_date": f"{y}-12-31", "upload_locked": False, "active": True, "created_at": now_iso()})

@app.on_event("startup")
async def startup():
    await seed()
    try:
        init_storage(); logger.info("Storage initialized")
    except Exception as e:
        logger.error(f"Storage init failed: {e}")

app.include_router(api_router)
app.add_middleware(CORSMiddleware, allow_credentials=True,
                   allow_origins=[os.environ.get("FRONTEND_URL", "http://localhost:3000")],
                   allow_methods=["*"], allow_headers=["*"])
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
